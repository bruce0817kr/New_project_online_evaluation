"""
Evaluation Endpoints - 평가 관련 API
"""
from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func
from uuid import UUID
from typing import Dict, Any, List, Optional
from datetime import datetime
from pydantic import BaseModel
import io
import csv

from app.core.database import get_db
from app.core.security import get_current_user, require_admin
from app.models.user import User
from app.models.evaluation import Evaluation
from app.models.company import Company
from app.models.project import Project

router = APIRouter()


# ============== Schemas ==============

class EvaluationResponse(BaseModel):
    """평가 응답"""
    id: str
    project_id: str = None
    project_name: str = None
    company_id: str
    company_name: str
    evaluator_id: str
    order: float = None
    scores_data: dict = {}
    total_score: float = None
    overall_comment: str = None
    is_submitted: bool
    submitted_at: datetime = None
    created_at: datetime
    updated_at: datetime = None

    # 동적 배점표 정보
    scoring_template: dict = None

    class Config:
        orm_mode = True


class EvaluationUpdate(BaseModel):
    """평가 업데이트"""
    scores_data: dict = None
    overall_comment: str = None


class EvaluationSubmit(BaseModel):
    """평가 제출"""
    signature_data: str = None


# ============== Endpoints ==============

@router.get("/my", response_model=List[EvaluationResponse])
async def get_my_evaluations(
    status: Optional[str] = None,  # 'in_progress' | 'submitted'
    project_id: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    내 평가 목록 조회 (심사위원용)
    """
    query = db.query(Evaluation).filter(
        Evaluation.evaluator_id == current_user.id
    ).options(
        joinedload(Evaluation.company),
        joinedload(Evaluation.company).joinedload(Company.project)
    )

    # 상태 필터
    if status == "in_progress":
        query = query.filter(Evaluation.is_submitted == False)
    elif status == "submitted":
        query = query.filter(Evaluation.is_submitted == True)

    # 프로젝트 필터
    if project_id:
        query = query.filter(Evaluation.project_id == project_id)

    # 순서대로 정렬
    evaluations = query.order_by(Evaluation.order.asc().nullslast()).all()

    return [
        EvaluationResponse(
            id=str(evaluation.id),
            project_id=str(evaluation.project_id) if evaluation.project_id else None,
            project_name=evaluation.company.project.name if evaluation.company and evaluation.company.project else None,
            company_id=str(evaluation.company_id),
            company_name=evaluation.company.name if evaluation.company else "알 수 없음",
            evaluator_id=str(evaluation.evaluator_id),
            order=evaluation.order,
            scores_data=evaluation.scores_data or {},
            total_score=evaluation.total_score,
            overall_comment=evaluation.overall_comment,
            is_submitted=evaluation.is_submitted,
            submitted_at=evaluation.submitted_at,
            created_at=evaluation.created_at,
            updated_at=evaluation.updated_at
        )
        for evaluation in evaluations
    ]


@router.get("/{evaluation_id}", response_model=EvaluationResponse)
async def get_evaluation_detail(
    evaluation_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    평가 상세 조회 (배점표 템플릿 포함)
    """
    from app.models.scoring_template import ScoringTemplate

    evaluation = db.query(Evaluation).filter(
        Evaluation.id == evaluation_id
    ).options(
        joinedload(Evaluation.company),
        joinedload(Evaluation.company).joinedload(Company.project)
    ).first()

    if not evaluation:
        raise HTTPException(status_code=404, detail="평가를 찾을 수 없습니다")

    # 권한 체크: 본인 평가이거나 관리자
    if current_user.role != "admin" and str(evaluation.evaluator_id) != str(current_user.id):
        raise HTTPException(status_code=403, detail="접근 권한이 없습니다")

    # 배점표 템플릿 조회
    scoring_template = None
    if evaluation.company and evaluation.company.project and evaluation.company.project.scoring_template_id:
        template = db.query(ScoringTemplate).filter(
            ScoringTemplate.id == evaluation.company.project.scoring_template_id
        ).first()
        if template:
            scoring_template = {
                "id": str(template.id),
                "name": template.name,
                "total_score": template.total_score,
                "sections": template.sections.get("sections", []) if template.sections else []
            }

    return EvaluationResponse(
        id=str(evaluation.id),
        project_id=str(evaluation.project_id) if evaluation.project_id else None,
        project_name=evaluation.company.project.name if evaluation.company and evaluation.company.project else None,
        company_id=str(evaluation.company_id),
        company_name=evaluation.company.name if evaluation.company else "알 수 없음",
        evaluator_id=str(evaluation.evaluator_id),
        order=evaluation.order,
        scores_data=evaluation.scores_data or {},
        total_score=evaluation.total_score,
        overall_comment=evaluation.overall_comment,
        is_submitted=evaluation.is_submitted,
        submitted_at=evaluation.submitted_at,
        created_at=evaluation.created_at,
        updated_at=evaluation.updated_at,
        scoring_template=scoring_template
    )


@router.put("/{evaluation_id}", response_model=EvaluationResponse)
async def update_evaluation(
    evaluation_id: str,
    data: EvaluationUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    평가 임시 저장
    """
    evaluation = db.query(Evaluation).filter(
        Evaluation.id == evaluation_id,
        Evaluation.evaluator_id == current_user.id
    ).first()

    if not evaluation:
        raise HTTPException(status_code=404, detail="평가를 찾을 수 없습니다")

    if evaluation.is_submitted:
        raise HTTPException(status_code=400, detail="이미 제출된 평가는 수정할 수 없습니다")

    # 업데이트
    if data.scores_data is not None:
        # 점수 변경 이력 기록
        from app.models.score_history import ScoreHistory

        old_scores = evaluation.scores_data or {}
        new_scores = data.scores_data

        # 변경된 항목 찾기 및 이력 기록
        for item_id, new_score in new_scores.items():
            if isinstance(new_score, (int, float)):
                old_score = old_scores.get(item_id)

                # 값이 변경되었거나 새로 추가된 경우
                if old_score != new_score:
                    history = ScoreHistory(
                        evaluation_id=evaluation.id,
                        evaluator_id=current_user.id,
                        item_id=item_id,
                        old_score=float(old_score) if old_score is not None else None,
                        new_score=float(new_score),
                        change_type="CREATE" if old_score is None else "UPDATE"
                    )
                    db.add(history)

        evaluation.scores_data = data.scores_data
        # 총점 계산 (간단 버전)
        if data.scores_data:
            total = sum([v for v in data.scores_data.values() if isinstance(v, (int, float))])
            evaluation.total_score = total

    if data.overall_comment is not None:
        evaluation.overall_comment = data.overall_comment

    evaluation.updated_at = datetime.utcnow()

    db.commit()
    db.refresh(evaluation)

    # 관계 로드
    db.refresh(evaluation, ["company"])
    if evaluation.company:
        db.refresh(evaluation.company, ["project"])

    return EvaluationResponse(
        id=str(evaluation.id),
        project_id=str(evaluation.project_id) if evaluation.project_id else None,
        project_name=evaluation.company.project.name if evaluation.company and evaluation.company.project else None,
        company_id=str(evaluation.company_id),
        company_name=evaluation.company.name if evaluation.company else "알 수 없음",
        evaluator_id=str(evaluation.evaluator_id),
        order=evaluation.order,
        scores_data=evaluation.scores_data or {},
        total_score=evaluation.total_score,
        overall_comment=evaluation.overall_comment,
        is_submitted=evaluation.is_submitted,
        submitted_at=evaluation.submitted_at,
        created_at=evaluation.created_at,
        updated_at=evaluation.updated_at
    )


@router.post("/{evaluation_id}/submit", response_model=EvaluationResponse)
async def submit_evaluation(
    evaluation_id: str,
    data: EvaluationSubmit = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    평가 최종 제출
    """
    evaluation = db.query(Evaluation).filter(
        Evaluation.id == evaluation_id,
        Evaluation.evaluator_id == current_user.id
    ).first()

    if not evaluation:
        raise HTTPException(status_code=404, detail="평가를 찾을 수 없습니다")

    if evaluation.is_submitted:
        raise HTTPException(status_code=400, detail="이미 제출된 평가입니다")

    # 제출 처리
    evaluation.is_submitted = True
    evaluation.submitted_at = datetime.utcnow()

    if data and data.signature_data:
        evaluation.signature_data = data.signature_data

    db.commit()
    db.refresh(evaluation)

    # 관계 로드
    db.refresh(evaluation, ["company"])
    if evaluation.company:
        db.refresh(evaluation.company, ["project"])

    return EvaluationResponse(
        id=str(evaluation.id),
        project_id=str(evaluation.project_id) if evaluation.project_id else None,
        project_name=evaluation.company.project.name if evaluation.company and evaluation.company.project else None,
        company_id=str(evaluation.company_id),
        company_name=evaluation.company.name if evaluation.company else "알 수 없음",
        evaluator_id=str(evaluation.evaluator_id),
        order=evaluation.order,
        scores_data=evaluation.scores_data or {},
        total_score=evaluation.total_score,
        overall_comment=evaluation.overall_comment,
        is_submitted=evaluation.is_submitted,
        submitted_at=evaluation.submitted_at,
        created_at=evaluation.created_at,
        updated_at=evaluation.updated_at
    )


@router.get("/project/{project_id}", response_model=List[EvaluationResponse])
async def get_evaluations_by_project(
    project_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin)
):
    """
    프로젝트별 평가 목록 조회 (관리자)
    """
    evaluations = db.query(Evaluation).filter(
        Evaluation.project_id == project_id
    ).options(
        joinedload(Evaluation.company),
        joinedload(Evaluation.evaluator)
    ).all()

    return [
        EvaluationResponse(
            id=str(evaluation.id),
            project_id=str(evaluation.project_id) if evaluation.project_id else None,
            project_name=evaluation.company.project.name if evaluation.company and evaluation.company.project else None,
            company_id=str(evaluation.company_id),
            company_name=evaluation.company.name if evaluation.company else "알 수 없음",
            evaluator_id=str(evaluation.evaluator_id),
            order=evaluation.order,
            scores_data=evaluation.scores_data or {},
            total_score=evaluation.total_score,
            overall_comment=evaluation.overall_comment,
            is_submitted=evaluation.is_submitted,
            submitted_at=evaluation.submitted_at,
            created_at=evaluation.created_at,
            updated_at=evaluation.updated_at
        )
        for evaluation in evaluations
    ]


@router.get("/company/{company_id}", response_model=List[EvaluationResponse])
async def get_evaluations_by_company(
    company_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin)
):
    """
    기업별 평가 목록 조회 (관리자)
    """
    evaluations = db.query(Evaluation).filter(
        Evaluation.company_id == company_id
    ).options(
        joinedload(Evaluation.company),
        joinedload(Evaluation.evaluator)
    ).all()

    return [
        EvaluationResponse(
            id=str(evaluation.id),
            project_id=str(evaluation.project_id) if evaluation.project_id else None,
            project_name=evaluation.company.project.name if evaluation.company and evaluation.company.project else None,
            company_id=str(evaluation.company_id),
            company_name=evaluation.company.name if evaluation.company else "알 수 없음",
            evaluator_id=str(evaluation.evaluator_id),
            order=evaluation.order,
            scores_data=evaluation.scores_data or {},
            total_score=evaluation.total_score,
            overall_comment=evaluation.overall_comment,
            is_submitted=evaluation.is_submitted,
            submitted_at=evaluation.submitted_at,
            created_at=evaluation.created_at,
            updated_at=evaluation.updated_at
        )
        for evaluation in evaluations
    ]


@router.get("/company/{company_id}/aggregate")
async def get_aggregated_scores(
    company_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin)
):
    """
    기업별 전체 심사위원 점수 집계
    최고/최저점 제외 평균 계산
    """
    evaluations = db.query(Evaluation).filter(
        Evaluation.company_id == company_id,
        Evaluation.is_submitted == True
    ).all()

    if not evaluations:
        return {
            "company_id": company_id,
            "average_score": 0.0,
            "trimmed_average": 0.0,
            "evaluator_count": 0
        }

    scores = [e.total_score for e in evaluations if e.total_score is not None]

    if not scores:
        return {
            "company_id": company_id,
            "average_score": 0.0,
            "trimmed_average": 0.0,
            "evaluator_count": len(evaluations)
        }

    # 평균 계산
    average_score = sum(scores) / len(scores)

    # 최고/최저 제외 평균 (3명 이상일 때만)
    trimmed_average = average_score
    if len(scores) >= 3:
        sorted_scores = sorted(scores)
        trimmed_scores = sorted_scores[1:-1]  # 최고, 최저 제외
        trimmed_average = sum(trimmed_scores) / len(trimmed_scores)

    return {
        "company_id": company_id,
        "average_score": round(average_score, 2),
        "trimmed_average": round(trimmed_average, 2),
        "evaluator_count": len(evaluations),
        "submitted_count": len(scores)
    }


@router.get("/project/{project_id}/rankings")
async def get_project_rankings(
    project_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin)
):
    """
    프로젝트별 기업 평가 결과 랭킹 조회
    - 최고/최저점 제외 평균 계산
    - 보너스/가산점 적용 (향후 확장)
    - 랭킹 부여
    """
    from app.models.company import Company
    from app.models.project import Project

    # 프로젝트 확인
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="프로젝트를 찾을 수 없습니다")

    # 해당 프로젝트의 모든 기업 조회
    companies = db.query(Company).filter(Company.project_id == project_id).all()

    results = []
    for company in companies:
        # 제출된 평가 조회
        evaluations = db.query(Evaluation).filter(
            Evaluation.company_id == company.id,
            Evaluation.is_submitted == True
        ).all()

        if not evaluations:
            results.append({
                "company_id": str(company.id),
                "company_name": company.name,
                "average_score": 0.0,
                "trimmed_average": 0.0,
                "bonus_score": 0.0,
                "final_score": 0.0,
                "evaluator_count": 0,
                "submitted_count": 0,
                "rank": None
            })
            continue

        scores = [e.total_score for e in evaluations if e.total_score is not None]

        if not scores:
            results.append({
                "company_id": str(company.id),
                "company_name": company.name,
                "average_score": 0.0,
                "trimmed_average": 0.0,
                "bonus_score": 0.0,
                "final_score": 0.0,
                "evaluator_count": len(evaluations),
                "submitted_count": 0,
                "rank": None
            })
            continue

        # 평균 계산
        average_score = sum(scores) / len(scores)

        # 최고/최저 제외 평균 (3명 이상일 때만)
        trimmed_average = average_score
        if len(scores) >= 3:
            sorted_scores = sorted(scores)
            trimmed_scores = sorted_scores[1:-1]
            trimmed_average = sum(trimmed_scores) / len(trimmed_scores)

        # 보너스 점수 (company.metadata에서 가져오기)
        bonus_score = 0.0
        if company.metadata and "bonus_score" in company.metadata:
            bonus_score = float(company.metadata.get("bonus_score", 0.0))

        # 최종 점수
        final_score = trimmed_average + bonus_score

        results.append({
            "company_id": str(company.id),
            "company_name": company.name,
            "average_score": round(average_score, 2),
            "trimmed_average": round(trimmed_average, 2),
            "bonus_score": round(bonus_score, 2),
            "final_score": round(final_score, 2),
            "evaluator_count": len(evaluations),
            "submitted_count": len(scores)
        })

    # 최종 점수 기준 정렬 및 랭킹 부여
    results.sort(key=lambda x: x["final_score"], reverse=True)

    for idx, result in enumerate(results):
        if result["final_score"] > 0:
            result["rank"] = idx + 1
        else:
            result["rank"] = None

    return {
        "project_id": project_id,
        "project_name": project.name,
        "total_companies": len(companies),
        "rankings": results
    }


@router.patch("/company/{company_id}/bonus")
async def update_company_bonus(
    company_id: str,
    data: dict,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin)
):
    """
    기업별 보너스 점수 설정
    (중소기업 가점, 지역 가점 등)
    """
    from app.models.company import Company

    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        raise HTTPException(status_code=404, detail="기업을 찾을 수 없습니다")

    bonus_score = data.get("bonus_score", 0.0)
    bonus_reason = data.get("bonus_reason", "")

    # metadata JSON 필드에 보너스 정보 저장
    metadata = company.metadata or {}
    metadata["bonus_score"] = bonus_score
    metadata["bonus_reason"] = bonus_reason

    company.metadata = metadata
    db.commit()

    return {
        "message": "보너스 점수가 설정되었습니다",
        "company_id": company_id,
        "bonus_score": bonus_score,
        "bonus_reason": bonus_reason
    }


@router.get("/project/{project_id}/export-csv")
async def export_rankings_csv(
    project_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin)
):
    """
    프로젝트별 평가 결과 CSV 내보내기
    """
    # rankings 데이터 조회 (기존 rankings 엔드포인트 로직 재사용)
    from app.models.company import Company
    from app.models.project import Project

    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="프로젝트를 찾을 수 없습니다")

    companies = db.query(Company).filter(Company.project_id == project_id).all()

    results = []
    for company in companies:
        evaluations = db.query(Evaluation).filter(
            Evaluation.company_id == company.id,
            Evaluation.is_submitted == True
        ).all()

        if not evaluations:
            continue

        scores = [e.total_score for e in evaluations if e.total_score is not None]
        if not scores:
            continue

        average_score = sum(scores) / len(scores)
        trimmed_average = average_score
        if len(scores) >= 3:
            sorted_scores = sorted(scores)
            trimmed_scores = sorted_scores[1:-1]
            trimmed_average = sum(trimmed_scores) / len(trimmed_scores)

        bonus_score = 0.0
        if company.metadata and "bonus_score" in company.metadata:
            bonus_score = float(company.metadata.get("bonus_score", 0.0))

        final_score = trimmed_average + bonus_score

        results.append({
            "company_name": company.name,
            "average_score": round(average_score, 2),
            "trimmed_average": round(trimmed_average, 2),
            "bonus_score": round(bonus_score, 2),
            "final_score": round(final_score, 2),
            "evaluator_count": len(evaluations),
            "submitted_count": len(scores)
        })

    # 점수 기준 정렬
    results.sort(key=lambda x: x["final_score"], reverse=True)

    # CSV 생성
    output = io.StringIO()
    writer = csv.DictWriter(
        output,
        fieldnames=["순위", "기업명", "평균점수", "조정평균", "가산점", "최종점수", "평가위원수", "제출수"],
        extrasaction='ignore'
    )

    writer.writeheader()
    for idx, result in enumerate(results):
        writer.writerow({
            "순위": idx + 1,
            "기업명": result["company_name"],
            "평균점수": result["average_score"],
            "조정평균": result["trimmed_average"],
            "가산점": result["bonus_score"],
            "최종점수": result["final_score"],
            "평가위원수": result["evaluator_count"],
            "제출수": result["submitted_count"]
        })

    # CSV 파일로 반환
    output.seek(0)
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8-sig")),  # BOM 추가 (Excel 한글 깨짐 방지)
        media_type="text/csv",
        headers={
            "Content-Disposition": f"attachment; filename={project.name}_평가결과_{datetime.now().strftime('%Y%m%d')}.csv"
        }
    )


@router.get("/{evaluation_id}/history")
async def get_score_history(
    evaluation_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin)
):
    """
    평가별 점수 변경 이력 조회 (관리자)
    """
    from app.models.score_history import ScoreHistory

    evaluation = db.query(Evaluation).filter(Evaluation.id == evaluation_id).first()
    if not evaluation:
        raise HTTPException(status_code=404, detail="평가를 찾을 수 없습니다")

    history = db.query(ScoreHistory).filter(
        ScoreHistory.evaluation_id == evaluation_id
    ).order_by(ScoreHistory.created_at.desc()).all()

    return {
        "evaluation_id": evaluation_id,
        "total_changes": len(history),
        "history": [
            {
                "id": str(h.id),
                "item_id": h.item_id,
                "item_name": h.item_name,
                "old_score": h.old_score,
                "new_score": h.new_score,
                "change_type": h.change_type,
                "evaluator_id": str(h.evaluator_id) if h.evaluator_id else None,
                "created_at": h.created_at.isoformat() if h.created_at else None
            }
            for h in history
        ]
    }


@router.get("/project/{project_id}/export-excel")
async def export_rankings_excel(
    project_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin)
):
    """
    프로젝트별 평가 결과 Excel 내보내기
    openpyxl 필요: pip install openpyxl
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="openpyxl이 설치되지 않았습니다. pip install openpyxl을 실행하세요"
        )

    from app.models.company import Company
    from app.models.project import Project

    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="프로젝트를 찾을 수 없습니다")

    companies = db.query(Company).filter(Company.project_id == project_id).all()

    results = []
    for company in companies:
        evaluations = db.query(Evaluation).filter(
            Evaluation.company_id == company.id,
            Evaluation.is_submitted == True
        ).all()

        if not evaluations:
            continue

        scores = [e.total_score for e in evaluations if e.total_score is not None]
        if not scores:
            continue

        average_score = sum(scores) / len(scores)
        trimmed_average = average_score
        if len(scores) >= 3:
            sorted_scores = sorted(scores)
            trimmed_scores = sorted_scores[1:-1]
            trimmed_average = sum(trimmed_scores) / len(trimmed_scores)

        bonus_score = 0.0
        if company.metadata and "bonus_score" in company.metadata:
            bonus_score = float(company.metadata.get("bonus_score", 0.0))

        final_score = trimmed_average + bonus_score

        results.append({
            "company_name": company.name,
            "average_score": round(average_score, 2),
            "trimmed_average": round(trimmed_average, 2),
            "bonus_score": round(bonus_score, 2),
            "final_score": round(final_score, 2),
            "evaluator_count": len(evaluations),
            "submitted_count": len(scores)
        })

    results.sort(key=lambda x: x["final_score"], reverse=True)

    # Excel 워크북 생성
    wb = Workbook()
    ws = wb.active
    ws.title = "평가 결과"

    # 헤더 스타일
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center")

    # 헤더 작성
    headers = ["순위", "기업명", "평균점수", "조정평균", "가산점", "최종점수", "평가위원수", "제출수"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # 데이터 작성
    for row_idx, result in enumerate(results, start=2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1)
        ws.cell(row=row_idx, column=2, value=result["company_name"])
        ws.cell(row=row_idx, column=3, value=result["average_score"])
        ws.cell(row=row_idx, column=4, value=result["trimmed_average"])
        ws.cell(row=row_idx, column=5, value=result["bonus_score"])
        ws.cell(row=row_idx, column=6, value=result["final_score"])
        ws.cell(row=row_idx, column=7, value=result["evaluator_count"])
        ws.cell(row=row_idx, column=8, value=result["submitted_count"])

        # 점수 셀 가운데 정렬
        for col_idx in range(3, 9):
            ws.cell(row=row_idx, column=col_idx).alignment = Alignment(horizontal="center")

    # 열 너비 자동 조정
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2) * 1.2
        ws.column_dimensions[column].width = adjusted_width

    # 메모리에 저장
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={project.name}_평가결과_{datetime.now().strftime('%Y%m%d')}.xlsx"
        }
    )
